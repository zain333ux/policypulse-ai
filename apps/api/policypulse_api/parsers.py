from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

import fitz
import pandas as pd
from charset_normalizer import from_bytes
from docx import Document

COMMENT_COLUMNS = (
    "comment",
    "comments",
    "feedback",
    "text",
    "review",
    "reviews",
    "response",
    "responses",
    "content",
    "message",
    "suggestion",
    "open-ended feedback",
    "additional comments",
)
IGNORED_RESPONSE_COLUMNS = (
    "timestamp",
    "email address",
    "email",
    "name",
    "username",
    "respondent id",
)


@dataclass(slots=True)
class ParsedPolicy:
    text: str
    warnings: list[str] = field(default_factory=list)
    metadata: dict[str, str | int | bool] = field(default_factory=dict)


@dataclass(slots=True)
class ParsedComments:
    comments: list[str]
    warnings: list[str] = field(default_factory=list)
    metadata: dict[str, str | int | bool] = field(default_factory=dict)


def parse_policy_bytes(filename: str, content: bytes) -> ParsedPolicy:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(content)
    if suffix == ".docx":
        return _parse_docx(content)
    if suffix in {".txt", ".md", ".rtf"}:
        text, encoding = _decode_text(content)
        return ParsedPolicy(
            text=_clean_document_text(text),
            metadata={"format": suffix.removeprefix("."), "encoding": encoding},
        )
    raise ValueError("Policy files must be PDF, DOCX, TXT, MD, or RTF.")


def parse_comments_bytes(filename: str, content: bytes) -> ParsedComments:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        text, encoding = _decode_text(content)
        comments = parse_comments_text(text)
        return ParsedComments(
            comments=comments,
            metadata={"format": suffix.removeprefix("."), "encoding": encoding, "rows": len(comments)},
        )
    if suffix == ".csv":
        return _parse_tabular_comments(filename, content, "csv")
    if suffix in {".xlsx", ".xls"}:
        return _parse_tabular_comments(filename, content, "excel")
    raise ValueError("Feedback files must be CSV, XLSX, XLS, TXT, or MD.")


def parse_comments_text(raw_text: str) -> list[str]:
    lines = [re.sub(r"^[\s\-•*\d.)]+", "", line).strip() for line in raw_text.splitlines()]
    return _deduplicate_comments([line for line in lines if line])


def _parse_pdf(content: bytes) -> ParsedPolicy:
    warnings: list[str] = []
    pages: list[str] = []
    with fitz.open(stream=content, filetype="pdf") as document:
        if document.needs_pass:
            raise ValueError("The PDF is password protected. Remove the password and upload it again.")
        for index, page in enumerate(document, start=1):
            text = page.get_text("text", sort=True).strip()
            if not text:
                try:
                    text_page = page.get_textpage_ocr(flags=0, dpi=200, full=True)
                    text = page.get_text("text", textpage=text_page, sort=True).strip()
                    if text:
                        warnings.append(f"Page {index} required OCR.")
                except RuntimeError:
                    warnings.append(
                        f"Page {index} contains no extractable text. Install Tesseract OCR "
                        "on the API host to analyze scanned pages."
                    )
            if text:
                pages.append(f"[Page {index}]\n{text}")
        if not pages:
            raise ValueError("No readable text was found in the PDF. It may be scanned, image-only, or corrupted.")
        return ParsedPolicy(
            text=_clean_document_text("\n\n".join(pages)),
            warnings=warnings,
            metadata={
                "format": "pdf",
                "pages": document.page_count,
                "ocr_used": any("required OCR" in warning for warning in warnings),
            },
        )


def _parse_docx(content: bytes) -> ParsedPolicy:
    document = Document(BytesIO(content))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    table_rows = 0
    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            values = [_clean_cell(cell.text) for cell in row.cells]
            if any(values):
                rows.append(" | ".join(values))
                table_rows += 1
        if rows:
            blocks.append(f"[Table {table_index}]\n" + "\n".join(rows))
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            if paragraph.text.strip():
                blocks.insert(0, f"[Header] {paragraph.text.strip()}")
        for paragraph in section.footer.paragraphs:
            if paragraph.text.strip():
                blocks.append(f"[Footer] {paragraph.text.strip()}")
    text = _clean_document_text("\n\n".join(blocks))
    if not text:
        raise ValueError("The DOCX file contains no readable paragraphs or tables.")
    return ParsedPolicy(
        text=text,
        metadata={
            "format": "docx",
            "paragraphs": len(document.paragraphs),
            "tables": len(document.tables),
            "table_rows": table_rows,
        },
    )


def _parse_tabular_comments(filename: str, content: bytes, kind: str) -> ParsedComments:
    warnings: list[str] = []
    if kind == "csv":
        decoded, encoding = _decode_text(content)
        delimiter = _detect_delimiter(decoded)
        frame = pd.read_csv(StringIO(decoded), sep=delimiter, dtype=str, keep_default_na=False)
        metadata: dict[str, str | int | bool] = {
            "format": "csv",
            "encoding": encoding,
            "delimiter": delimiter,
        }
    else:
        engine = "openpyxl" if Path(filename).suffix.lower() == ".xlsx" else None
        workbook = pd.read_excel(BytesIO(content), sheet_name=None, dtype=str, engine=engine)
        non_empty = {name: sheet for name, sheet in workbook.items() if not sheet.empty}
        if not non_empty:
            raise ValueError("The Excel workbook contains no response rows.")
        frame = pd.concat(
            [sheet.assign(__sheet=name) for name, sheet in non_empty.items()],
            ignore_index=True,
        )
        metadata = {"format": "excel", "sheets": len(non_empty)}

    if frame.empty:
        raise ValueError("The feedback file contains no response rows.")
    frame.columns = [str(column).strip() for column in frame.columns]
    comments, selected_columns = _comments_from_frame(frame)
    if not comments:
        raise ValueError(
            "No analyzable feedback text was found. Include a comment/feedback column "
            "or open-ended Google Forms responses."
        )
    if len(selected_columns) > 1:
        warnings.append("Multiple response columns were combined per respondent to preserve Google Forms answers.")
    metadata.update(
        {
            "rows": len(frame.index),
            "comments": len(comments),
            "columns_used": ", ".join(selected_columns),
        }
    )
    return ParsedComments(comments=comments, warnings=warnings, metadata=metadata)


def _comments_from_frame(frame: pd.DataFrame) -> tuple[list[str], list[str]]:
    normalized = {_normalize_column(column): column for column in frame.columns}
    direct = [normalized[name] for name in COMMENT_COLUMNS if name in normalized]
    if direct:
        selected = list(dict.fromkeys(direct))
    else:
        selected = [
            column
            for column in frame.columns
            if _normalize_column(column) not in IGNORED_RESPONSE_COLUMNS
            and column != "__sheet"
            and _is_textual_response_column(frame[column])
        ]
    comments: list[str] = []
    for _, row in frame.iterrows():
        parts: list[str] = []
        for column in selected:
            value = _clean_cell(row.get(column, ""))
            if not value or value.lower() in {"nan", "none", "n/a"}:
                continue
            if len(selected) == 1:
                parts.append(value)
            else:
                parts.append(f"{column}: {value}")
        if parts:
            comments.append(" | ".join(parts))
    return _deduplicate_comments(comments), [str(column) for column in selected]


def _is_textual_response_column(series: pd.Series) -> bool:
    values = [str(value).strip() for value in series if str(value).strip()]
    if not values:
        return False
    average_length = sum(len(value) for value in values) / len(values)
    unique_ratio = len(set(values)) / len(values)
    return average_length >= 12 or unique_ratio >= 0.45


def _decode_text(content: bytes) -> tuple[str, str]:
    match = from_bytes(content).best()
    if not match:
        raise ValueError("The text encoding could not be detected.")
    return str(match), match.encoding or "unknown"


def _detect_delimiter(text: str) -> str:
    try:
        return csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _normalize_column(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value).strip().lower())


def _clean_cell(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value)).strip()


def _clean_document_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _deduplicate_comments(comments: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for comment in comments:
        key = re.sub(r"\W+", "", comment).lower()
        if key and key not in seen:
            seen.add(key)
            result.append(comment.strip())
    return result
