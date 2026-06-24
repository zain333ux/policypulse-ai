from io import BytesIO

from docx import Document
from openpyxl import Workbook

from policypulse_api.parsers import parse_comments_bytes, parse_comments_text, parse_policy_bytes


def test_parse_comments_text_removes_empty_lines_and_bullets() -> None:
    assert parse_comments_text("- First comment\n\n• Second comment") == [
        "First comment",
        "Second comment",
    ]


def test_parse_comments_csv_detects_feedback_column() -> None:
    content = b"id,feedback\n1,Clear appeal rules are needed\n2,Attendance matters\n"
    parsed = parse_comments_bytes("comments.csv", content)
    assert parsed.comments == [
        "Clear appeal rules are needed",
        "Attendance matters",
    ]
    assert parsed.metadata["columns_used"] == "feedback"


def test_parse_policy_txt_handles_utf8_bom() -> None:
    parsed = parse_policy_bytes("policy.txt", b"\xef\xbb\xbfA clear policy statement.")
    assert parsed.text == "A clear policy statement."
    assert parsed.metadata["encoding"] == "utf_8"


def test_parse_google_forms_excel_combines_open_response_columns() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Form Responses 1"
    sheet.append(["Timestamp", "Do you support the policy?", "What should improve?"])
    sheet.append(["2026-06-24", "Neutral", "Add a documented appeal process."])
    stream = BytesIO()
    workbook.save(stream)

    parsed = parse_comments_bytes("responses.xlsx", stream.getvalue())

    assert parsed.comments == [
        "Do you support the policy?: Neutral | What should improve?: Add a documented appeal process."
    ]
    assert parsed.metadata["format"] == "excel"


def test_parse_docx_includes_table_content() -> None:
    document = Document()
    document.add_paragraph("Attendance Policy")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Threshold"
    table.cell(0, 1).text = "85%"
    table.cell(1, 0).text = "Penalty"
    table.cell(1, 1).text = "Exam restriction"
    stream = BytesIO()
    document.save(stream)

    parsed = parse_policy_bytes("policy.docx", stream.getvalue())

    assert "Threshold | 85%" in parsed.text
    assert parsed.metadata["tables"] == 1
